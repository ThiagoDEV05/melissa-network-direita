#!/usr/bin/env python3
"""
Networking da Direita: servidor local do formulário de cadastro.

Serve a página da pasta site/ e grava cada cadastro em
dados/cadastros-networking.xlsx (com cópia de segurança em CSV).
Nada sai deste computador.

Uso:  python3 servidor.py            (computador + celulares no mesmo Wi-Fi)
      python3 servidor.py --so-local (só neste computador)
"""
import csv
import json
import os
import re
import socket
import sys
import threading
import webbrowser
from datetime import datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("Falta o pacote openpyxl. Instale com:  pip3 install openpyxl")

BASE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.join(BASE, "site")
DADOS = os.path.join(BASE, "dados")
PLANILHA = os.path.join(DADOS, "cadastros-networking.xlsx")
BACKUP_CSV = os.path.join(DADOS, "backup-cadastros.csv")
PORTA = int(os.environ.get("PORTA", "8080"))
SO_LOCAL = "--so-local" in sys.argv

COLUNAS = [
    ("ID", 11), ("Data do cadastro", 17), ("Status", 12), ("Nome completo", 28),
    ("Nome da empresa", 24), ("Profissão", 20), ("Categoria profissional", 28),
    ("Especialidade", 22), ("Cidade e estado", 20), ("Instagram", 20),
    ("Link do Instagram", 34), ("WhatsApp comercial", 18), ("Descrição dos serviços", 50),
    ("Interesse em parcerias", 14), ("Autorizou divulgação", 12), ("Observações da administração", 30),
]
CABECALHO = [c for c, _ in COLUNAS]
CATEGORIAS = [
    "Marketing e Comunicação", "Advocacia e Jurídico", "Saúde e Bem-estar", "Beleza e Estética",
    "Contabilidade e Finanças", "Tecnologia e Informática", "Vendas e Comércio",
    "Empreendedorismo e Consultoria", "Imóveis e Construção", "Educação e Cursos",
    "Alimentação e Gastronomia", "Prestadores de Serviços", "Outros Profissionais",
]
PARCERIAS = ["Sim", "A combinar", "Não"]
STATUS = ["Novo", "Aprovado", "Recusado", "Inativo"]
# False: o cadastro aparece na lista do site na hora (Recusado ou Inativo tira do ar).
# True: só aparece depois que o Status for mudado para Aprovado.
EXIGIR_APROVACAO = False

trava = threading.Lock()
fila = []  # cadastros que não couberam na planilha (ex.: arquivo aberto e travado)


def criar_planilha():
    os.makedirs(DADOS, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "CADASTROS"
    verde = PatternFill("solid", start_color="0A6B3B")
    for i, (nome, largura) in enumerate(COLUNAS, start=1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = verde
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = largura
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUNAS)).column_letter}2000"

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    dv.add("C2:C2000")
    ws.add_data_validation(dv)
    cores = {"Novo": "FFF5D2", "Aprovado": "E4EFE7", "Recusado": "F6E1E1", "Inativo": "EEEEEE"}
    for st, cor in cores.items():
        ws.conditional_formatting.add(
            "A2:P2000", FormulaRule(formula=[f'$C2="{st}"'], fill=PatternFill("solid", start_color=cor, end_color=cor)))

    leg = wb.create_sheet("COMO USAR")
    leg.column_dimensions["A"].width = 100
    linhas = [
        "NETWORKING DA DIREITA: planilha de cadastros",
        "",
        "Cada envio do formulário vira uma linha na aba CADASTROS com Status \"Novo\" (amarelo)",
        "e já aparece na lista de profissionais do site.",
        "Para tirar alguém do site, mude o Status para Recusado ou Inativo (lista suspensa).",
        "Use as setinhas do cabeçalho para filtrar, por exemplo, só os Novos ou uma categoria.",
        "A coluna \"Observações da administração\" é livre para anotações.",
        "",
        "IMPORTANTE: feche esta planilha no Excel/Numbers enquanto o formulário estiver recebendo cadastros.",
        "Se ela estiver aberta e você salvar por cima, cadastros novos podem sumir daqui.",
        "Mesmo assim nada se perde: todo envio também vai para dados/backup-cadastros.csv.",
    ]
    for i, t in enumerate(linhas, start=1):
        c = leg.cell(row=i, column=1, value=t)
        c.font = Font(name="Arial", bold=(i == 1), size=14 if i == 1 else 11, color="0A6B3B" if i == 1 else "222222")
    wb.save(PLANILHA)


def texto(v, maximo):
    return re.sub(r"[\r\n\t]+", " ", str(v or "")).strip()[:maximo]


def validar(d):
    faltando = [n for k, n in [("nome", "nome"), ("profissao", "profissão"), ("categoria", "categoria"),
                               ("cidade", "cidade"), ("instagram", "Instagram"), ("descricao", "descrição"),
                               ("parcerias", "interesse em parcerias")] if not texto(d.get(k), 300)]
    if faltando:
        raise ValueError("Preencha: " + ", ".join(faltando) + ".")
    if d.get("autorizacao") is not True:
        raise ValueError("É preciso autorizar a divulgação dos dados.")
    usuario = re.sub(r"^https?://(www\.)?instagram\.com/", "", texto(d["instagram"], 100).lower()).lstrip("@").split("/")[0]
    usuario = re.sub(r"[^a-z0-9._]", "", usuario)[:30]
    if not usuario:
        raise ValueError("Instagram inválido.")
    whats = re.sub(r"\D", "", str(d.get("whatsapp") or ""))[:13]
    return {
        "Nome completo": texto(d["nome"], 80),
        "Nome da empresa": texto(d.get("empresa"), 80),
        "Profissão": texto(d["profissao"], 60),
        "Categoria profissional": d["categoria"] if d["categoria"] in CATEGORIAS else "Outros Profissionais",
        "Especialidade": texto(d.get("especialidade"), 60),
        "Cidade e estado": texto(d["cidade"], 60),
        "Instagram": "@" + usuario,
        "Link do Instagram": f"https://www.instagram.com/{usuario}/",
        "WhatsApp comercial": whats,
        "Descrição dos serviços": texto(d["descricao"], 280),
        "Interesse em parcerias": d["parcerias"] if d["parcerias"] in PARCERIAS else "A combinar",
        "Autorizou divulgação": "Sim",
        "Observações da administração": "",
    }


def proximo_id(ws):
    maior = 0
    for (v,) in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        m = re.search(r"(\d+)", str(v or ""))
        if m:
            maior = max(maior, int(m.group(1)))
    return maior + 1


def gravar_csv(reg):
    novo = not os.path.exists(BACKUP_CSV)
    with open(BACKUP_CSV, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=CABECALHO, delimiter=";")
        if novo:
            w.writeheader()
        linha = {}
        for k, v in reg.items():
            if isinstance(v, datetime):
                v = v.strftime("%d/%m/%Y %H:%M")
            elif isinstance(v, str) and v[:1] in "=+-@":
                v = "'" + v  # impede que o Excel trate o texto como fórmula
            linha[k] = v
        w.writerow(linha)


def gravar_planilha(registros):
    """Acrescenta os registros na planilha. Lança exceção se o arquivo estiver travado."""
    if not os.path.exists(PLANILHA):
        criar_planilha()
    wb = load_workbook(PLANILHA)
    ws = wb["CADASTROS"]
    ultima = ws.max_row
    while ultima > 1 and all(c.value in (None, "") for c in ws[ultima]):
        ultima -= 1
    for reg in registros:
        ultima += 1
        for col, nome in enumerate(CABECALHO, start=1):
            c = ws.cell(row=ultima, column=col, value=reg.get(nome, ""))
            c.font = Font(name="Arial")
            if isinstance(c.value, str):
                c.data_type = "s"  # texto vindo do formulário nunca vira fórmula
            if nome == "Data do cadastro":
                c.number_format = "dd/mm/yyyy hh:mm"
            elif nome == "Link do Instagram" and reg.get(nome):
                c.hyperlink = reg[nome]
                c.font = Font(name="Arial", color="1B3A8C", underline="single")
    tmp = PLANILHA + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, PLANILHA)


def ler_registros():
    """Linhas da planilha (mais as que ainda estão na fila) como dicionários."""
    regs = []
    if os.path.exists(PLANILHA):
        wb = load_workbook(PLANILHA, read_only=True)
        linhas = wb["CADASTROS"].iter_rows(values_only=True)
        cab = [str(c or "").strip() for c in next(linhas, ())]
        regs = [dict(zip(cab, l)) for l in linhas if any(v not in (None, "") for v in l)]
        wb.close()
    return regs + list(fila)


def publico(r):
    """Cartão público de um cadastro, ou None se não deve aparecer no site."""
    s = lambda k: str(r.get(k) or "").strip()
    st = s("Status")
    if st in ("Recusado", "Inativo") or (EXIGIR_APROVACAO and st != "Aprovado"):
        return None
    if s("Autorizou divulgação") != "Sim" or not s("Nome completo"):
        return None
    return {
        "id": s("ID"), "nome": s("Nome completo"), "empresa": s("Nome da empresa"),
        "profissao": s("Profissão"), "especialidade": s("Especialidade"),
        "categoria": s("Categoria profissional"), "cidade": s("Cidade e estado"),
        "instagram": s("Instagram").lstrip("@"), "whatsapp": re.sub(r"\D", "", s("WhatsApp comercial")),
        "descricao": s("Descrição dos serviços"),
    }


def listar():
    with trava:
        return [p for p in map(publico, ler_registros()) if p]


def salvar_cadastro(d):
    campos = validar(d)
    with trava:
        if not os.path.exists(PLANILHA):
            criar_planilha()
        try:
            existentes = ler_registros()
        except Exception:
            existentes = list(fila)
        if any(str(r.get("Instagram") or "").strip().lower() == campos["Instagram"]
               and str(r.get("Status") or "").strip() != "Recusado" for r in existentes):
            raise ValueError("Este Instagram já está cadastrado no guia.")
        try:
            n = proximo_id(load_workbook(PLANILHA, read_only=True)["CADASTROS"])
        except Exception:
            n = 1
        n = max([n] + [int(r["ID"][4:]) + 1 for r in fila])
        reg = {"ID": f"NDD-{n:04d}", "Data do cadastro": datetime.now().replace(microsecond=0), "Status": "Novo", **campos}
        gravar_csv(reg)
        fila.append(reg)
        try:
            gravar_planilha(fila)
            fila.clear()
            print(f"  ✔ {reg['ID']}  {reg['Nome completo']} ({reg['Profissão']}, {reg['Cidade e estado']})")
        except Exception as e:
            print(f"  ! {reg['ID']} salvo só no backup CSV: a planilha parece aberta ou travada ({e}). "
                  "Feche a planilha; o próximo envio grava os pendentes.")
    perfil = publico(reg)
    return {"id": reg["ID"], "publicado": perfil is not None, "perfil": perfil}


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **k):
        super().__init__(*a, directory=SITE, **k)

    def log_message(self, *_):
        pass

    def end_headers(self):
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Cache-Control", "no-store")
        super().end_headers()

    def json(self, codigo, obj):
        corpo = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(codigo)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(corpo)))
        self.end_headers()
        self.wfile.write(corpo)

    def do_GET(self):
        if self.path.split("?")[0] == "/api/profissionais":
            try:
                return self.json(200, {"ok": True, "profissionais": listar()})
            except Exception as e:
                print("  ! não consegui ler a planilha:", e)
                return self.json(500, {"ok": False, "erro": "Não foi possível ler a planilha."})
        return super().do_GET()

    def do_POST(self):
        if self.path.split("?")[0] != "/api/cadastro":
            return self.json(404, {"ok": False, "erro": "Endereço não encontrado."})
        try:
            tam = int(self.headers.get("Content-Length") or 0)
            if tam <= 0 or tam > 20000:
                raise ValueError("Envio vazio ou grande demais.")
            d = json.loads(self.rfile.read(tam).decode("utf-8"))
            if not isinstance(d, dict):
                raise ValueError("Formato inválido.")
            if d.get("website"):  # campo invisível: só robôs preenchem
                return self.json(200, {"ok": True})
            return self.json(200, {"ok": True, **salvar_cadastro(d)})
        except ValueError as e:
            return self.json(400, {"ok": False, "erro": str(e)})
        except Exception as e:
            print("  ! erro inesperado:", e)
            return self.json(500, {"ok": False, "erro": "Erro ao salvar. Tente de novo."})


def ip_da_rede():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("10.255.255.255", 1))  # não envia nada; só descobre o IP local
        return s.getsockname()[0]
    except OSError:
        return None
    finally:
        s.close()


def main():
    if not os.path.exists(PLANILHA):
        criar_planilha()
    host = "127.0.0.1" if SO_LOCAL else "0.0.0.0"
    try:
        srv = ThreadingHTTPServer((host, PORTA), Handler)
    except OSError:
        sys.exit(f"A porta {PORTA} já está em uso. O servidor já está aberto em outra janela? "
                 f"Ou rode com outra porta:  PORTA=8090 python3 servidor.py")
    url = f"http://localhost:{PORTA}"
    print("\n  NETWORKING DA DIREITA: formulário no ar\n")
    print(f"  Neste computador:   {url}")
    ip = None if SO_LOCAL else ip_da_rede()
    if ip:
        print(f"  Celular no Wi-Fi:   http://{ip}:{PORTA}")
    print(f"  Planilha:           {PLANILHA}")
    print("\n  Deixe esta janela aberta. Para desligar: Ctrl+C ou feche a janela.\n")
    threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\n  Servidor desligado.")


if __name__ == "__main__":
    main()
